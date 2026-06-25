"""Utility functions — diagram layout, helpers, etc."""

from __future__ import annotations

import base64
import io
import logging
import math
from typing import Any, List

import httpx
from docx import Document
from langchain_core.messages import AnyMessage, HumanMessage

from agent.models import (
    DiagramEdge,
    DiagramModel,
    DiagramNode,
    Position,
    RelationType,
    SchemaModel,
    Viewport,
)

logger = logging.getLogger(__name__)

# ── Async image pre-fetch ────────────────────────────────────────────────────


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _extract_docx_text(content: bytes) -> str:
    """Extract plain text from a DOCX file bytes."""
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


async def resolve_image_urls(messages: List[AnyMessage]) -> List[AnyMessage]:
    """Replace HTTPS image_url parts with inline base64 data URLs.

    langchain_google_genai fetches HTTP image URLs synchronously via requests,
    which triggers a BlockingError inside LangGraph's async event loop.
    This function pre-fetches all remote images asynchronously with httpx and
    converts them to data URIs so the LLM library never makes a blocking call.

    DOCX files are not supported by Gemini as image_url — they are downloaded
    and converted to plain text parts instead.
    """
    resolved: List[AnyMessage] = []
    async with httpx.AsyncClient(timeout=30) as client:
        for msg in messages:
            if not isinstance(msg, HumanMessage):
                resolved.append(msg)
                continue

            content = msg.content
            if not isinstance(content, list):
                resolved.append(msg)
                continue

            new_parts: List[Any] = []
            changed = False
            for part in content:
                if (
                    isinstance(part, dict)
                    and part.get("type") == "image_url"
                    and isinstance(part.get("image_url"), dict)
                ):
                    url: str = part["image_url"].get("url", "")
                    if url.startswith("http://") or url.startswith("https://"):
                        try:
                            resp = await client.get(url)
                            resp.raise_for_status()
                            mime = resp.headers.get("content-type", "image/jpeg").split(";")[0]

                            if mime == DOCX_MIME or url.lower().endswith(".docx"):
                                # Gemini doesn't support DOCX — extract text instead
                                text = _extract_docx_text(resp.content)
                                new_parts.append({"type": "text", "text": f"[Word document content]\n{text}"})
                                changed = True
                                logger.debug("Extracted text from DOCX URL: %s", url[:80])
                            else:
                                b64 = base64.b64encode(resp.content).decode()
                                data_url = f"data:{mime};base64,{b64}"
                                new_parts.append({"type": "image_url", "image_url": {"url": data_url}})
                                changed = True
                                logger.debug("Resolved image URL to base64: %s", url[:80])
                            continue
                        except Exception as exc:
                            logger.warning("Failed to fetch URL %s: %s", url[:80], exc)
                            # Replace with text placeholder so the LLM library never
                            # sees the raw HTTPS URL — it would fetch it synchronously
                            # via requests, triggering a BlockingError in the ASGI loop.
                            new_parts.append({"type": "text", "text": f"[Image unavailable: {url[:80]}]"})
                            changed = True
                            continue
                new_parts.append(part)

            if changed:
                resolved.append(HumanMessage(content=new_parts))
            else:
                resolved.append(msg)

    return resolved

# Layout constants
NODE_WIDTH = 220
FIELD_ROW_HEIGHT = 28
NODE_HEADER_HEIGHT = 44
NODE_PADDING = 16
GAP_X = 300
GAP_Y = 60


def _compute_node_height(field_count: int) -> float:
    """Calculate the pixel height of a table node based on its field count."""
    return NODE_HEADER_HEIGHT + NODE_PADDING + max(field_count, 1) * FIELD_ROW_HEIGHT


def _cardinality_label(rel_type: RelationType) -> str:
    """Return a human-readable label for a relationship type."""
    return {
        RelationType.ONE_TO_ONE: "1:1",
        RelationType.ONE_TO_MANY: "1:N",
        RelationType.MANY_TO_MANY: "N:N",
    }.get(rel_type, "")


def compute_diagram_layout(schema: SchemaModel) -> DiagramModel:
    """Compute a grid-based auto-layout for the given schema.

    Tables are arranged in a grid with ``cols`` columns, distributing
    them as evenly as possible to keep the layout roughly square.
    """
    tables = schema.tables
    n = len(tables)
    if n == 0:
        return DiagramModel()

    cols = max(1, math.ceil(math.sqrt(n)))

    nodes: list[DiagramNode] = []
    for idx, table in enumerate(tables):
        row, col = divmod(idx, cols)
        height = _compute_node_height(len(table.fields))
        nodes.append(
            DiagramNode(
                id=table.id,
                position=Position(
                    x=col * (NODE_WIDTH + GAP_X),
                    y=row * (height + GAP_Y),
                ),
                width=NODE_WIDTH,
                height=height,
            )
        )

    edges: list[DiagramEdge] = []
    for rel in schema.relationships:
        edges.append(
            DiagramEdge(
                id=rel.id,
                source=rel.source_table_id,
                target=rel.target_table_id,
                source_handle=rel.source_field_id,
                target_handle=rel.target_field_id,
                label=_cardinality_label(rel.type),
            )
        )

    return DiagramModel(
        nodes=nodes,
        edges=edges,
        viewport=Viewport(x=0, y=0, zoom=1.0),
    )
